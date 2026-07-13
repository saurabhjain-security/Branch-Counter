#!/usr/bin/env python3
"""
Branch Counter
--------------
Script  : GitHub Repository Branch Counter
Author  : Saurabh Jain
Version : 1.0

A tool to count branches across all repositories in one or more GitHub
organisations. Outputs a rich terminal dashboard plus optional export to
Excel, CSV, or JSON.

Pre-requisite:
  pip install rich requests openpyxl

  Token (one of):
    export GITHUB_TOKEN=ghp_...
    export GH_TOKEN=ghp_...
    gh auth login   (fallback)

Usage:
  python branch_counter.py
  python branch_counter.py --org carter-rmn shyftlabs
  python branch_counter.py --min-branches 5
  python branch_counter.py --output-format csv
  python branch_counter.py --exclude-archived --exclude-forks
  python branch_counter.py --sort count
  python branch_counter.py --filter bloated
  python branch_counter.py --verbose
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import signal
import subprocess
import sys
import time
from concurrent.futures import Future, ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from threading import Lock, local as thread_local
from typing import Dict, List, Optional, Set, Tuple, TypedDict

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rich import box
from rich.console import Console
from rich.panel import Panel
from rich.progress import (
    BarColumn,
    MofNCompleteColumn,
    Progress,
    SpinnerColumn,
    TaskProgressColumn,
    TextColumn,
    TimeElapsedColumn,
)
from rich.rule import Rule
from rich.table import Table
from rich.traceback import install as install_rich_traceback

install_rich_traceback(show_locals=False)

# ── Constants ─────────────────────────────────────────────────────────────────
GITHUB_API   = "https://api.github.com"
VERSION      = "1.0"
AUTHOR       = "Saurabh Jain"
MAX_RETRIES  = 5
RL_THROTTLE  = 200    # slow down when remaining calls drop below this

DEFAULT_ORGS = ["carter-rmn", "shyftlabs"]

# Branch-count thresholds for colour coding
THRESH_LOW    = 15    # ≤ this → green  (healthy)
THRESH_MID    = 50    # ≤ this → yellow (moderate)
THRESH_HIGH   = 70    # ≤ this → orange (crowded)
# > THRESH_HIGH → red (bloated)

console    = Console()
print_lock = Lock()
_shutdown  = False


# ── TypedDict schema ──────────────────────────────────────────────────────────

class RepoRow(TypedDict):
    org_name:     str
    repo_name:    str
    repo_url:     str
    is_fork:      bool
    is_archived:  bool
    branch_count: int
    health:       str   # "Healthy" | "Moderate" | "Crowded" | "Bloated" | "Archived"


# ── Excel styling ──────────────────────────────────────────────────────────────
_HDR_FILL        = PatternFill("solid", start_color="1F3864")
_HDR_FONT        = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_BODY_FONT       = Font(name="Calibri", size=10)
_THIN            = Side(style="thin", color="D9D9D9")
_BORDER          = Border(top=_THIN, bottom=_THIN, left=_THIN, right=_THIN)
_FILL_HEALTHY    = PatternFill("solid", start_color="C6EFCE")   # green
_FILL_MODERATE   = PatternFill("solid", start_color="FFEB9C")   # yellow
_FILL_CROWDED    = PatternFill("solid", start_color="FFD8A8")   # orange
_FILL_BLOATED    = PatternFill("solid", start_color="FFC7CE")   # red
_FILL_ARCHIVED   = PatternFill("solid", start_color="D9D9D9")   # grey
_FILL_NONE       = PatternFill("solid", start_color="FFFFFF")

VERBOSE = False


# ── Logging helpers ───────────────────────────────────────────────────────────

def vlog(msg: str) -> None:
    if not VERBOSE:
        return
    ts = datetime.now(timezone.utc).strftime("%H:%M:%S")
    with print_lock:
        console.print(f"  [dim cyan]VRB {ts}[/]  {msg}")


# ── Banner ────────────────────────────────────────────────────────────────────

def print_banner(args: argparse.Namespace) -> None:
    art = (
        " ██████╗ ██████╗  █████╗ ███╗   ██╗ ██████╗██╗  ██╗\n"
        " ██╔══██╗██╔══██╗██╔══██╗████╗  ██║██╔════╝██║  ██║\n"
        " ██████╔╝██████╔╝███████║██╔██╗ ██║██║     ███████║\n"
        " ██╔══██╗██╔══██╗██╔══██║██║╚██╗██║██║     ██╔══██║\n"
        " ██████╔╝██║  ██║██║  ██║██║ ╚████║╚██████╗██║  ██║\n"
        " ╚═════╝ ╚═╝  ╚═╝╚═╝  ╚═╝╚═╝  ╚═══╝ ╚═════╝╚═╝  ╚═╝\n"
        "\n"
        " ██████╗ ██████╗ ██╗   ██╗███╗   ██╗████████╗███████╗██████╗ \n"
        " ██╔════╝██╔═══██╗██║   ██║████╗  ██║╚══██╔══╝██╔════╝██╔══██╗\n"
        " ██║     ██║   ██║██║   ██║██╔██╗ ██║   ██║   █████╗  ██████╔╝\n"
        " ██║     ██║   ██║██║   ██║██║╚██╗██║   ██║   ██╔══╝  ██╔══██╗\n"
        " ╚██████╗╚██████╔╝╚██████╔╝██║ ╚████║   ██║   ███████╗██║  ██║\n"
        "  ╚═════╝ ╚═════╝  ╚═════╝ ╚═╝  ╚═══╝   ╚═╝   ╚══════╝╚═╝  ╚═╝"
    )
    flags = []
    if args.exclude_forks:    flags.append("forks excluded")
    if args.exclude_archived: flags.append("archived excluded")
    if args.min_branches > 0: flags.append(f"min branches: {args.min_branches}")
    if VERBOSE:               flags.append("[bold cyan]VERBOSE[/]")
    flags_str = "  [dim]·[/]  ".join(flags) if flags else ""

    console.print()
    console.print(Panel(
        f"[bold bright_cyan]{art}[/]\n\n"
        f"  [bold bright_green]GitHub Organisation Repository Branch Counter[/]"
        f"   [dim]|[/]   [dim]v{VERSION}[/]"
        f"   [dim]|[/]   [bold yellow]Author: {AUTHOR}[/]\n\n"
        f"  [dim cyan]Healthy threshold  :[/] [bold white]≤ {THRESH_LOW} branches[/]\n"
        f"  [dim cyan]Moderate threshold :[/] [bold white]{THRESH_LOW+1}–{THRESH_MID} branches[/]\n"
        f"  [dim cyan]Crowded threshold  :[/] [bold white]{THRESH_MID+1}–{THRESH_HIGH} branches[/]\n"
        f"  [dim cyan]Bloated threshold  :[/] [bold white]> {THRESH_HIGH} branches[/]\n"

        f"  [dim cyan]Workers            :[/] [bold white]{args.workers}[/]"
        + (f"\n  [dim cyan]Flags              :[/] {flags_str}" if flags_str else ""),
        border_style="bright_blue",
        box=box.DOUBLE_EDGE,
        expand=False,
        padding=(1, 2),
    ))
    console.print()


# ── CLI ───────────────────────────────────────────────────────────────────────

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        prog="branch_counter",
        description=f"GitHub Repository Branch Counter v{VERSION}",
        epilog=(
            "Examples:\n"
            "  python branch_counter.py\n"
            "  python branch_counter.py --org acme-corp\n"
            "  python branch_counter.py --min-branches 10\n"
            "  python branch_counter.py --filter bloated --output-format csv\n"
            "  python branch_counter.py --exclude-archived --sort count\n"
            "  python branch_counter.py --verbose\n"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    p.add_argument("--org", nargs="+", metavar="ORG", default=None,
                   help=f"Target org(s). Default: {', '.join(DEFAULT_ORGS)}.")
    p.add_argument("--min-branches", dest="min_branches", type=int, default=0, metavar="N",
                   help="Only show repos with at least N branches (default: 0 = all).")
    p.add_argument("--output", "-o", default="",
                   help="Output file path (default: branch_counter_<timestamp>.<ext>).")
    p.add_argument("--output-format", dest="output_format",
                   choices=["xlsx", "csv", "json"], default="xlsx",
                   help="Export format: xlsx (default) | csv | json.")
    p.add_argument("--no-export", dest="no_export", action="store_true",
                   help="Skip file export — terminal output only.")
    p.add_argument("--exclude-forks", dest="exclude_forks", action="store_true",
                   help="Skip forked repositories.")
    p.add_argument("--exclude-archived", dest="exclude_archived", action="store_true",
                   help="Skip archived repositories entirely.")
    p.add_argument("--filter", dest="filter_mode",
                   choices=["all", "healthy", "moderate", "crowded", "bloated", "archived"],
                   default="all",
                   help="Display subset: all (default) | healthy | moderate | crowded | bloated | archived.")
    p.add_argument("--sort", dest="sort_by",
                   choices=["count", "name", "org", "health"], default="count",
                   help="Sort results by: count (default, desc) | name | org | health.")
    p.add_argument("--workers", type=int, default=10, metavar="N",
                   help="Parallel workers (default: 10).")
    p.add_argument("--verbose", "-v", action="store_true", default=False,
                   help="Show per-call API logging.")
    return p.parse_args()


# ── Auth ──────────────────────────────────────────────────────────────────────

def get_token() -> str:
    for env_var in ("GITHUB_TOKEN", "GH_TOKEN"):
        token = os.environ.get(env_var, "").strip()
        if token:
            console.print(f"  [green]✓[/]  Token from env var [bold cyan]{env_var}[/]")
            return token
    try:
        r = subprocess.run(["gh", "auth", "token"],
                           capture_output=True, text=True, timeout=10)
    except FileNotFoundError:
        console.print(Panel(
            "[bold red]✗  No token found[/]\n\n"
            "  Set [bold cyan]GITHUB_TOKEN[/] or [bold cyan]GH_TOKEN[/] env var,\n"
            "  or install gh CLI and run [bold cyan]gh auth login[/]",
            border_style="red", box=box.ROUNDED, expand=False))
        sys.exit(1)
    except subprocess.TimeoutExpired:
        console.print("  [bold red]✗[/]  `gh auth token` timed out.")
        sys.exit(1)

    token = r.stdout.strip()
    if not token or r.returncode != 0:
        console.print(Panel(
            "[bold red]✗  Not authenticated[/]\n\n"
            "  Set [bold cyan]GITHUB_TOKEN[/] env var, or run [bold cyan]gh auth login[/]",
            border_style="red", box=box.ROUNDED, expand=False))
        sys.exit(1)
    console.print("  [green]✓[/]  Token via [bold cyan]gh auth token[/]")
    return token


# ── GitHub API Client ─────────────────────────────────────────────────────────

_tl = thread_local()


def _make_session(token: str) -> requests.Session:
    s = requests.Session()
    s.headers.update({
        "Authorization":        f"Bearer {token}",
        "Accept":               "application/vnd.github+json",
        "X-GitHub-Api-Version": "2022-11-28",
    })
    return s


class GHClient:
    def __init__(self, token: str) -> None:
        self._token        = token
        self._rl_remaining = 5000
        self._rl_limit     = 5000
        self._rl_lock      = Lock()

    def _session(self) -> requests.Session:
        if not hasattr(_tl, "session"):
            _tl.session = _make_session(self._token)
        return _tl.session

    def _get(self, path: str, params: Optional[Dict] = None) -> requests.Response:
        url     = f"{GITHUB_API}{path}" if path.startswith("/") else path
        backoff = 2
        attempt = 0

        while attempt < MAX_RETRIES:
            with self._rl_lock:
                remaining = self._rl_remaining
            if remaining < RL_THROTTLE and remaining > 0:
                vlog(f"RL low ({remaining}) — pausing 2s")
                time.sleep(2)

            try:
                t0   = time.monotonic()
                resp = self._session().get(url, params=params, timeout=30)
                ms   = (time.monotonic() - t0) * 1000
            except (requests.ConnectionError, requests.Timeout) as exc:
                attempt += 1
                vlog(f"Network error ({exc.__class__.__name__}) attempt {attempt}/{MAX_RETRIES}")
                time.sleep(backoff)
                backoff = min(backoff * 2, 60)
                continue

            try:
                with self._rl_lock:
                    self._rl_remaining = int(
                        resp.headers.get("X-RateLimit-Remaining", self._rl_remaining))
                    self._rl_limit = int(
                        resp.headers.get("X-RateLimit-Limit", self._rl_limit))
            except (ValueError, TypeError):
                pass

            vlog(
                f"[dim]GET[/] {url[:80]}  "
                f"[{'bold green' if resp.ok else 'bold red'}]{resp.status_code}[/]  "
                f"[dim]{ms:.0f}ms  RL={self._rl_remaining}[/]"
            )

            is_rate_limited = (
                resp.status_code == 429
                or (resp.status_code == 403
                    and resp.headers.get("X-RateLimit-Remaining") == "0")
            )
            if is_rate_limited:
                reset     = int(resp.headers.get("X-RateLimit-Reset", time.time() + 60))
                wait      = max(reset - int(time.time()), 1)
                reset_utc = datetime.fromtimestamp(reset, tz=timezone.utc).strftime("%H:%M:%S UTC")
                with print_lock:
                    console.print(
                        f"  [bold yellow]⚠[/]  Rate limit — "
                        f"waiting [bold]{wait}s[/]  [dim](resets {reset_utc})[/]"
                    )
                time.sleep(wait + 1)
                continue

            if resp.status_code == 403:
                vlog(f"[bold red]403 permission denied[/] {url[:80]}")
                return resp

            if resp.status_code >= 500:
                attempt += 1
                vlog(f"Server error {resp.status_code} — retry {attempt}/{MAX_RETRIES}")
                time.sleep(backoff)
                backoff = min(backoff * 2, 60)
                continue

            return resp

        raise RuntimeError(f"Gave up after {MAX_RETRIES} attempts: {url}")

    def paginate(self, path: str, params: Optional[Dict] = None) -> Tuple[List[dict], bool]:
        p = dict(params or {})
        p.setdefault("per_page", 100)
        p["page"] = 1
        collected: List[dict] = []
        while True:
            try:
                resp = self._get(path, params=p)
            except RuntimeError as exc:
                with print_lock:
                    console.print(f"  [bold red]✗[/]  {exc}")
                return collected, False
            if not resp.ok:
                return collected, False
            items = resp.json()
            if not isinstance(items, list) or not items:
                return collected, True
            collected.extend(items)
            if 'rel="next"' not in resp.headers.get("Link", ""):
                return collected, True
            p["page"] += 1

    def get_one(self, path: str) -> Optional[dict]:
        try:
            resp = self._get(path)
            return resp.json() if resp.ok else None
        except RuntimeError:
            return None

    def rate_summary(self) -> str:
        with self._rl_lock:
            return f"{self._rl_remaining:,}/{self._rl_limit:,}"

    def check_rate_limit(self) -> None:
        data = self.get_one("/rate_limit")
        if not data:
            return
        rate      = data.get("rate", {})
        remaining = rate.get("remaining", 0)
        limit     = rate.get("limit", 5000)
        reset_ts  = rate.get("reset", 0)
        reset_utc = datetime.fromtimestamp(reset_ts, tz=timezone.utc).strftime("%H:%M:%S UTC")
        with self._rl_lock:
            self._rl_remaining = remaining
            self._rl_limit     = limit
        colour = "green" if remaining > 1000 else ("yellow" if remaining > 200 else "red")
        console.print(
            f"  [green]✓[/]  Rate limit: "
            f"[bold {colour}]{remaining:,}[/] / {limit:,} calls remaining  "
            f"[dim](resets {reset_utc})[/]"
        )
        if remaining < 100:
            console.print(
                "  [bold red]⚠  Fewer than 100 API calls remaining.[/]  "
                "Results may be incomplete."
            )


# ── Helpers ───────────────────────────────────────────────────────────────────

def classify(count: int, is_archived: bool) -> str:
    if is_archived:
        return "Archived"
    if count <= THRESH_LOW:
        return "Healthy"
    if count <= THRESH_MID:
        return "Moderate"
    if count <= THRESH_HIGH:
        return "Crowded"
    return "Bloated"


def count_colour(count: int, is_archived: bool) -> str:
    if is_archived:
        return f"[dim white]{count}[/]"
    if count <= THRESH_LOW:
        return f"[bold green]{count}[/]"
    if count <= THRESH_MID:
        return f"[bold yellow]{count}[/]"
    if count <= THRESH_HIGH:
        return f"[bold orange3]{count}[/]"
    return f"[bold red]{count}[/]"


def health_colour(health: str) -> str:
    m = {
        "Healthy":  "[bold green]Healthy[/]",
        "Moderate": "[bold yellow]Moderate[/]",
        "Crowded":  "[bold orange3]Crowded[/]",
        "Bloated":  "[bold red]Bloated[/]",
        "Archived": "[dim white]Archived[/]",
    }
    return m.get(health, health)


def _c(val: int, color: str) -> str:
    return f"[{color}]{val}[/]" if val else "[dim]0[/]"


# ── Data collection ───────────────────────────────────────────────────────────

def get_branch_count(client: GHClient, full_name: str) -> int:
    branches, _ = client.paginate(f"/repos/{full_name}/branches")
    return len(branches)


def fetch_repo_row(
    client: GHClient,
    org: str,
    repo: dict,
) -> RepoRow:
    if _shutdown:
        raise InterruptedError("Shutdown requested")

    repo_name   = repo.get("name", "")
    full_name   = repo.get("full_name", "")
    html_url    = repo.get("html_url") or f"https://github.com/{full_name}"
    is_archived = bool(repo.get("archived", False))
    is_fork     = bool(repo.get("fork", False))

    count  = get_branch_count(client, full_name)
    health = classify(count, is_archived)

    return RepoRow(
        org_name=org,
        repo_name=repo_name,
        repo_url=html_url,
        is_fork=is_fork,
        is_archived=is_archived,
        branch_count=count,
        health=health,
    )


def fetch_org_repos(
    client: GHClient,
    org: str,
    exclude_forks: bool,
    exclude_archived: bool,
) -> List[dict]:
    repos, _ = client.paginate(f"/orgs/{org}/repos", params={"type": "all"})
    out = []
    for r in repos:
        if exclude_forks and r.get("fork"):
            continue
        if exclude_archived and r.get("archived"):
            continue
        out.append(r)
    return out


# ── Sorting & filtering ───────────────────────────────────────────────────────

_HEALTH_ORDER = {"Bloated": 0, "Crowded": 1, "Moderate": 2, "Healthy": 3, "Archived": 4}


def sort_rows(rows: List[RepoRow], sort_by: str) -> List[RepoRow]:
    if sort_by == "name":
        return sorted(rows, key=lambda r: r["repo_name"].lower())
    if sort_by == "org":
        return sorted(rows, key=lambda r: (r["org_name"].lower(), r["repo_name"].lower()))
    if sort_by == "health":
        return sorted(rows, key=lambda r: (
            _HEALTH_ORDER.get(r["health"], 9), -r["branch_count"]
        ))
    # default: count descending
    return sorted(rows, key=lambda r: -r["branch_count"])


def apply_filter(rows: List[RepoRow], mode: str, min_branches: int) -> List[RepoRow]:
    if mode != "all":
        rows = [r for r in rows if r["health"].lower() == mode]
    if min_branches > 0:
        rows = [r for r in rows if r["branch_count"] >= min_branches]
    return rows


# ── Terminal table ────────────────────────────────────────────────────────────

def print_results_table(rows: List[RepoRow], filter_mode: str, min_branches: int) -> None:
    display = apply_filter(rows, filter_mode, min_branches)

    console.print(Rule("[bold bright_cyan]  Scan Results  [/]", style="bright_blue"))
    console.print()

    filter_label = "" if filter_mode == "all" else f" — filter: [bold]{filter_mode}[/]"
    tbl = Table(
        title=(
            f"[bold bright_cyan]GitHub Repository Branch Analysis[/]  "
            f"[dim]({len(display)} repos{filter_label})[/]"
        ),
        box=box.ROUNDED,
        border_style="bright_blue",
        header_style="bold bright_blue",
        show_lines=True,
        expand=True,
    )
    tbl.add_column("Sno",          justify="right",  style="dim white",  width=5,    no_wrap=True)
    tbl.add_column("Org / Owner",  justify="left",   style="bold cyan",  min_width=16, no_wrap=True)
    tbl.add_column("Repo Name",    justify="left",   style="white",      min_width=20)
    tbl.add_column("Repo URL",     justify="left",   overflow="fold",    min_width=34)
    tbl.add_column("Fork",         justify="center", style="white",      width=7,    no_wrap=True)
    tbl.add_column("Archived",     justify="center", style="white",      width=10,   no_wrap=True)
    tbl.add_column("Branches",     justify="center", style="white",      width=10,   no_wrap=True)
    tbl.add_column("Health",       justify="center", style="white",      width=12,   no_wrap=True)

    for idx, row in enumerate(display, start=1):
        is_bloated      = row["health"] == "Bloated"
        url_markup      = (
            f"[bold red]{row['repo_url']}[/]" if is_bloated
            else f"[cyan]{row['repo_url']}[/]"
        )
        archived_markup = "[bold yellow]True[/]"  if row["is_archived"] else "[dim white]False[/]"
        fork_markup     = "[dim white]True[/]"    if row["is_fork"]     else "[dim white]False[/]"

        tbl.add_row(
            str(idx),
            row["org_name"],
            row["repo_name"],
            url_markup,
            fork_markup,
            archived_markup,
            count_colour(row["branch_count"], row["is_archived"]),
            health_colour(row["health"]),
        )

    console.print(tbl)
    console.print()


# ── Summary ───────────────────────────────────────────────────────────────────

def print_summary(
    rows: List[RepoRow],
    org_labels: List[str],
    elapsed: float,
    rate_summary_str: str,
) -> None:
    console.print(Rule("[bold bright_green]  Summary  [/]", style="bright_green"))
    console.print()

    per_org = Table(
        title="[bold white]Per-Organisation Breakdown[/]",
        box=box.ROUNDED, border_style="bright_green",
        header_style="bold bright_green", show_lines=True, expand=False,
    )
    per_org.add_column("Organisation", style="bold cyan", no_wrap=True, min_width=20)
    per_org.add_column("Repos",        justify="right",  style="white",  width=7)
    per_org.add_column("Branches",     justify="right",  width=9)
    per_org.add_column("Avg",          justify="right",  width=7)
    per_org.add_column("Max",          justify="right",  width=7)
    per_org.add_column("Healthy",      justify="right",  width=9)
    per_org.add_column("Moderate",     justify="right",  width=10)
    per_org.add_column("Crowded",      justify="right",  width=9)
    per_org.add_column("Bloated",      justify="right",  width=9)
    per_org.add_column("Archived",     justify="right",  width=9)
    per_org.add_column("Forks",        justify="right",  width=7)

    grand: Dict[str, int] = dict(
        repos=0, branches=0, max_b=0,
        healthy=0, moderate=0, crowded=0, bloated=0, archived=0, forks=0,
    )

    for org in org_labels:
        org_rows = [r for r in rows if r["org_name"] == org]
        if not org_rows:
            continue
        repos    = len(org_rows)
        branches = sum(r["branch_count"] for r in org_rows)
        max_b    = max(r["branch_count"] for r in org_rows)
        avg_b    = branches / repos if repos else 0
        healthy  = sum(1 for r in org_rows if r["health"] == "Healthy")
        moderate = sum(1 for r in org_rows if r["health"] == "Moderate")
        crowded  = sum(1 for r in org_rows if r["health"] == "Crowded")
        bloated  = sum(1 for r in org_rows if r["health"] == "Bloated")
        archived = sum(1 for r in org_rows if r["is_archived"])
        forks    = sum(1 for r in org_rows if r["is_fork"])

        grand["repos"]    += repos
        grand["branches"] += branches
        grand["max_b"]     = max(grand["max_b"], max_b)
        grand["healthy"]  += healthy
        grand["moderate"] += moderate
        grand["crowded"]  += crowded
        grand["bloated"]  += bloated
        grand["archived"] += archived
        grand["forks"]    += forks

        per_org.add_row(
            org,
            str(repos),
            str(branches),
            f"{avg_b:.1f}",
            str(max_b),
            _c(healthy,  "bold green"),
            _c(moderate, "bold yellow"),
            _c(crowded,  "bold orange3"),
            _c(bloated,  "bold red"),
            _c(archived, "dim white"),
            _c(forks,    "dim cyan"),
        )

    grand_repos = grand["repos"] or 1
    grand_avg   = grand["branches"] / grand_repos

    per_org.add_row(
        "[bold white]TOTAL[/]",
        f"[bold white]{grand['repos']}[/]",
        f"[bold white]{grand['branches']}[/]",
        f"[bold white]{grand_avg:.1f}[/]",
        f"[bold white]{grand['max_b']}[/]",
        _c(grand["healthy"],  "bold green"),
        _c(grand["moderate"], "bold yellow"),
        _c(grand["crowded"],  "bold orange3"),
        _c(grand["bloated"],  "bold red"),
        _c(grand["archived"], "dim white"),
        _c(grand["forks"],    "dim cyan"),
    )
    console.print(per_org)
    console.print()

    # Legend
    legend = Table(box=box.SIMPLE, border_style="bright_blue",
                   show_header=False, expand=False, padding=(0, 2))
    for _ in range(4):
        legend.add_column("x", justify="center")
    legend.add_row(
        f"[bold green]● Healthy[/]   [dim]≤ {THRESH_LOW} branches[/]",
        f"[bold yellow]● Moderate[/]  [dim]{THRESH_LOW+1}–{THRESH_MID} branches[/]",
        f"[bold orange3]● Crowded[/]   [dim]{THRESH_MID+1}–{THRESH_HIGH} branches[/]",
        f"[bold red]● Bloated[/]   [dim]> {THRESH_HIGH} branches[/]",
    )
    console.print(legend)
    console.print()

    # Stat cards
    cards = Table(box=box.SIMPLE_HEAVY, border_style="bright_blue",
                  show_header=False, expand=True, padding=(0, 2))
    for _ in range(6):
        cards.add_column("x", justify="center")

    def card(label: str, value: str, color: str) -> str:
        return f"[dim]{label}[/]\n[{color}]{value}[/{color}]"

    cards.add_row(
        card("Total repos",    str(grand["repos"]),    "bold white"),
        card("Total branches", str(grand["branches"]), "bold white"),
        card("Avg branches",   f"{grand_avg:.1f}",    "bold cyan"),
        card("Bloated repos",  str(grand["bloated"]), "bold red"),
        card("Healthy repos",  str(grand["healthy"]), "bold green"),
        card("API remaining",  rate_summary_str,       "dim cyan"),
    )
    console.print(cards)
    console.print()

    def _fmt(s: float) -> str:
        m, sec = divmod(int(s), 60)
        return f"{m}m {sec}s" if m else f"{sec:.0f}s"

    console.print(Panel(
        f"[bold green]✓  Analysis Complete[/]   [dim]·[/]   "
        f"[white]Orgs:[/] [bold]{len(org_labels)}[/]   "
        f"[white]Repos:[/] [bold]{grand['repos']}[/]   "
        f"[white]Total branches:[/] [bold]{grand['branches']}[/]   "
        f"[white]Avg:[/] [bold cyan]{grand_avg:.1f}[/]   "
        f"[white]Bloated:[/] [bold red]{grand['bloated']}[/]   "
        f"[white]Runtime:[/] [bold cyan]{_fmt(elapsed)}[/]",
        border_style="bright_green", box=box.DOUBLE_EDGE, expand=False,
    ))
    console.print()


# ── Export ────────────────────────────────────────────────────────────────────

COL_HEADERS = [
    "Sno", "Org / Owner", "Repo Name", "Repo URL",
    "Fork", "Archived", "Branch Count", "Health",
]

COL_WIDTHS = {
    "A": 6,   # Sno
    "B": 22,  # Org
    "C": 30,  # Repo Name
    "D": 55,  # URL
    "E": 8,   # Fork
    "F": 10,  # Archived
    "G": 14,  # Branch Count
    "H": 12,  # Health
}

_HEALTH_FILL = {
    "Healthy":  _FILL_HEALTHY,
    "Moderate": _FILL_MODERATE,
    "Crowded":  _FILL_CROWDED,
    "Bloated":  _FILL_BLOATED,
    "Archived": _FILL_ARCHIVED,
}


def _row_to_cells(i: int, row: RepoRow) -> List:
    return [
        i,
        row["org_name"],
        row["repo_name"],
        row["repo_url"],
        str(row["is_fork"]),
        str(row["is_archived"]),
        row["branch_count"],
        row["health"],
    ]


def _style_ws(ws) -> None:
    for col_letter, width in COL_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    for cell in ws[1]:
        cell.fill      = _HDR_FILL
        cell.font      = _HDR_FONT
        cell.border    = Border(
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")

    health_col_idx = COL_HEADERS.index("Health")

    for row_cells in ws.iter_rows(min_row=2, max_row=ws.max_row):
        health_val = str(row_cells[health_col_idx].value or "")
        fill       = _HEALTH_FILL.get(health_val, _FILL_NONE)
        border     = Border(
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
        )
        for ci, cell in enumerate(row_cells):
            cell.font      = _BODY_FONT
            cell.border    = border
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.fill      = fill if ci == health_col_idx else _FILL_NONE

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export_xlsx(rows: List[RepoRow], path: str, org_labels: List[str]) -> str:
    import re
    wb = Workbook()
    wb.remove(wb.active)

    def _write_sheet(ws, sheet_rows: List[RepoRow]) -> None:
        ws.append(COL_HEADERS)
        for i, row in enumerate(sheet_rows, start=1):
            ws.append(_row_to_cells(i, row))
        _style_ws(ws)

    ws_all = wb.create_sheet("All Repos")
    _write_sheet(ws_all, rows)

    for org in org_labels:
        org_rows = [r for r in rows if r["org_name"] == org]
        if not org_rows:
            continue
        sname  = re.sub(r'[\\/*?:\[\]]', '-', org)[:31]
        ws_org = wb.create_sheet(sname)
        _write_sheet(ws_org, org_rows)

    wb.save(path)
    return os.path.abspath(path)


def export_csv(rows: List[RepoRow], path: str) -> str:
    if not rows:
        console.print("  [dim yellow]⚠  No rows to export.[/]")
        return path
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(COL_HEADERS)
        for i, row in enumerate(rows, start=1):
            writer.writerow(_row_to_cells(i, row))
    return os.path.abspath(path)


def export_json(rows: List[RepoRow], path: str) -> str:
    out = []
    for i, row in enumerate(rows, start=1):
        out.append({
            "sno":          i,
            "org_name":     row["org_name"],
            "repo_name":    row["repo_name"],
            "repo_url":     row["repo_url"],
            "is_fork":      row["is_fork"],
            "is_archived":  row["is_archived"],
            "branch_count": row["branch_count"],
            "health":       row["health"],
        })
    with open(path, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
    return os.path.abspath(path)


# ── Progress bar ──────────────────────────────────────────────────────────────

def make_progress() -> Progress:
    return Progress(
        SpinnerColumn(spinner_name="dots", style="bright_blue"),
        TextColumn("[bold cyan]{task.description}"),
        BarColumn(bar_width=36, style="bright_blue", complete_style="bright_green"),
        TaskProgressColumn(),
        MofNCompleteColumn(),
        TimeElapsedColumn(),
        console=console,
        transient=False,
    )


# ── Main ──────────────────────────────────────────────────────────────────────

def main() -> None:
    global VERBOSE, _shutdown

    args    = parse_args()
    VERBOSE = args.verbose

    date_file = datetime.now().strftime("%Y-%m-%dT%H-%M")

    def _handle_sigint(sig, frame):
        global _shutdown
        _shutdown = True
        with print_lock:
            console.print("\n  [bold yellow]⚠  Interrupted — saving partial results...[/]")
    signal.signal(signal.SIGINT, _handle_sigint)

    t_start = time.monotonic()
    print_banner(args)

    # ── Auth ──────────────────────────────────────────────────────────
    with console.status("  [dim]Resolving GitHub token...[/]", spinner="dots"):
        token = get_token()

    client = GHClient(token)

    me = client.get_one("/user")
    if not me:
        console.print("[bold red]✗  Could not fetch user info — check token.[/]")
        sys.exit(1)

    console.print(
        f"  [green]✓[/]  Logged in as [bold cyan]{me.get('login')}[/]"
        f"  [dim]({me.get('name', '')})[/]"
    )
    console.print()
    client.check_rate_limit()
    console.print()

    # ── Orgs ──────────────────────────────────────────────────────────
    orgs = args.org if args.org else DEFAULT_ORGS
    console.print(
        f"  [green]✓[/]  Targeting [bold]{len(orgs)}[/] org(s): "
        f"[dim]{', '.join(orgs)}[/]"
    )
    console.print()

    # ── Collect repos concurrently ────────────────────────────────────
    repo_queue: List[Tuple[str, dict]] = []
    seen_full_names: Set[str] = set()

    with make_progress() as progress:
        task = progress.add_task("[cyan]Listing repositories...", total=len(orgs))
        with ThreadPoolExecutor(max_workers=min(len(orgs), 5)) as pool:
            fut_map: Dict[Future, str] = {
                pool.submit(
                    fetch_org_repos, client, org,
                    args.exclude_forks, args.exclude_archived,
                ): org
                for org in orgs
            }
            for fut in as_completed(fut_map):
                org = fut_map[fut]
                try:
                    for r in fut.result():
                        fn = r.get("full_name", "")
                        if fn not in seen_full_names:
                            seen_full_names.add(fn)
                            repo_queue.append((org, r))
                except Exception as exc:
                    with print_lock:
                        console.print(f"  [bold red]✗[/]  Listing {org}: {exc}")
                progress.advance(task)

    if not repo_queue:
        console.print(Panel(
            "[bold yellow]⚠  No repositories found to scan.[/]",
            border_style="yellow", box=box.ROUNDED, expand=False))
        sys.exit(0)

    console.print(
        f"  [green]✓[/]  [bold]{len(repo_queue)}[/] repo(s) queued"
        + (" (forks excluded)"    if args.exclude_forks    else "")
        + (" (archived excluded)" if args.exclude_archived else "")
    )
    console.print()

    # ── Concurrent branch count fetch ────────────────────────────────
    all_rows: List[RepoRow] = []

    with make_progress() as progress:
        task = progress.add_task("[cyan]Counting branches...", total=len(repo_queue))
        with ThreadPoolExecutor(max_workers=args.workers) as pool:
            future_map: Dict[Future, Tuple[str, str]] = {
                pool.submit(fetch_repo_row, client, org, repo): (org, repo.get("name", ""))
                for org, repo in repo_queue
            }
            for fut in as_completed(future_map):
                if _shutdown:
                    pool.shutdown(wait=False, cancel_futures=True)
                    break
                org_label, repo_name = future_map[fut]
                progress.update(
                    task,
                    description=f"[cyan]{org_label}[/] — [white]{repo_name[:28]}[/]",
                )
                try:
                    all_rows.append(fut.result())
                except InterruptedError:
                    pass
                except Exception as exc:
                    with print_lock:
                        console.print(
                            f"  [bold red]✗[/]  {org_label}/{repo_name}: {exc}"
                        )
                progress.advance(task)

    if not all_rows:
        console.print("[bold yellow]⚠  No results collected.[/]")
        sys.exit(0)

    all_rows = sort_rows(all_rows, args.sort_by)

    seen_orgs: Set[str] = set()
    org_labels: List[str] = []
    for r in all_rows:
        if r["org_name"] not in seen_orgs:
            seen_orgs.add(r["org_name"])
            org_labels.append(r["org_name"])

    # ── Terminal output ───────────────────────────────────────────────
    print_results_table(all_rows, args.filter_mode, args.min_branches)

    # ── Export ────────────────────────────────────────────────────────
    if not args.no_export:
        ext          = args.output_format
        default_name = f"branch_counter_{date_file}.{ext}"
        out_path     = args.output or default_name

        with console.status(f"[bright_blue]Writing {ext.upper()} report...[/]", spinner="dots"):
            if ext == "xlsx":
                saved = export_xlsx(all_rows, out_path, org_labels)
            elif ext == "csv":
                saved = export_csv(all_rows, out_path)
            else:
                saved = export_json(all_rows, out_path)

        console.print(f"  [green]✓[/]  {ext.upper()} saved → [dim]{saved}[/]")
        console.print()

    # ── Summary ───────────────────────────────────────────────────────
    elapsed = time.monotonic() - t_start
    print_summary(all_rows, org_labels, elapsed, client.rate_summary())

    if _shutdown:
        console.print("  [bold yellow]⚠  Scan was interrupted — output may be partial.[/]")
        console.print()


if __name__ == "__main__":
    main()